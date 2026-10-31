"""
Keeps the test plan workbook (test-cases.xlsx) in step with the QA suite.

  python3 scripts/qa/plan_workbook.py add-qa-cases          # the QA-automation module (24 Oct 2026)
  python3 scripts/qa/plan_workbook.py add-list-cases        # the Purchase lists module (24 Oct 2026)
  python3 scripts/qa/plan_workbook.py add-app-cases         # the App install & updates module (31 Oct 2026)
  python3 scripts/qa/plan_workbook.py statuses REGISTER.xlsx  # write each case's overall status back

add-qa-cases adds the cases and scenarios of the QA automation features
(QA-010 … QA-014) once — rows that already exist are left alone — and
extends the Run summary, Coverage and Change log sheets, the Status
drop-down, its colours and the filter to the new rows. add-list-cases does
the same for the Purchase lists module (PL-001 … PL-007, TC-LIST-001 … 028,
docs/purchase-lists-2026-10).

statuses copies the master register's Overall status and detail into the
plan's Status / Actual result / Defect / Run date columns (Pass (partial) is
written as Pass with the uncovered part in Actual result; Manual as Blocked),
so the workbook's own Run summary counts are live.

openpyxl keeps the workbook's formulas, styles and validations; run
scripts/qa/export-cases.ts afterwards so tests/qa/catalog matches.
"""
import re
import sys
from copy import copy

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList

PLAN = "test-cases.xlsx"
MODULE = "QA automation"

QA_CASES = [
    # id, features, title, type, priority, role, method, preconditions, steps, expected
    ("TC-QA-001", "QA-010, QA-003", "Every case of the plan is declared in the QA suite", "Positive", "P0", "All", "API",
     "tests/qa/catalog exported from this workbook",
     "1. List every TC and SC ID of the catalog\n2. Scan tests/qa/**/*.spec.ts for qa() and manual() declarations",
     "Every ID is declared at least once - automated as a check, or declared manual() with the reason. No ID of the plan is silently missing."),
    ("TC-QA-002", "QA-010", "Writing and local tiers never run on production", "Negative", "P0", "All", "API",
     "Suite pointed at https://gokesari.com",
     "1. Run the suite against gokesari.com\n2. Read the recorded rows of signed-in and local checks",
     "Every signed / local check is recorded 'skipped: not run on production'; the write helpers refuse production (assertWritable). Only GETs and requests the server must refuse reach gokesari.com."),
    ("TC-QA-003", "QA-010", "Each check records its case result", "Positive", "P1", "All", "API",
     "A finished run",
     "1. Read test-results/qa/<host>/*.jsonl\n2. Check one row per case per check",
     "Each row has the case ID, title, status (pass / fail / blocked / skipped), detail, tier, site, spec file and time; a failed or timed-out check is recorded as fail with its reason."),
    ("TC-QA-004", "QA-012", "Master register lists every case and scenario with a status per site", "Positive", "P0", "All", "API",
     "At least one collected run in docs/testing/qa-runs",
     "1. npx tsx scripts/qa/build-register.ts\n2. Open the workbook",
     "Test cases and E2E scenarios sheets hold every case and scenario of the plan once, each with Local, Staging, Production and Overall status, automation and spec file; Summary, Defects and Runs sheets are filled."),
    ("TC-QA-005", "QA-011", "Claude QA agent and skills are installed", "Positive", "P1", "All", "UI",
     "Repository checkout",
     "1. Open .claude/agents/gokesari-qa.md and .claude/skills/gokesari-qa-*/SKILL.md",
     "The agent and each skill have a name and description; the skills cover running the suite, signing the staging accounts in, writing a check and reporting."),
    ("TC-QA-006", "QA-013", "Scheduled QA workflow runs the production-safe tier on both sites", "Positive", "P1", "All", "UI",
     "Workflow .github/workflows/qa.yml on main",
     "1. Read the workflow's schedule and jobs\n2. Run it by hand once",
     "Scheduled weekly (Monday) and runnable by hand; runs tests/qa/guest against test.gokesari.com and gokesari.com and keeps the results as an artifact; never signs in or writes. By hand it can also run the whole suite on a server built from the commit."),
    ("TC-QA-007", "QA-014", "Local QA environment brings up the test data of the plan", "Positive", "P1", "Admin", "DB",
     "Local PostgreSQL and a built app",
     "1. bash scripts/qa/local-up.sh 3000 gk_qa\n2. Read test-results/qa/world-localhost_3000.json",
     "Database migrated and seeded; mail catcher and server up; the world file names every account of the Test data sheet (2 customers, 3 shop owners + a pending one, 2 riders, admin, operator, operator-limited, society admin), 4 shops and the listings the cases need."),
]

QA_SCENARIOS = [
    ("SC-36", "Release check on both sites", "QA-010, QA-012, QA-013", "P0",
     "1. The release reaches test.gokesari.com\n2. The production-safe and signed-in tiers run there; defects are fixed and re-deployed\n3. The release reaches gokesari.com\n4. The production-safe tier runs there\n5. The register is rebuilt",
     "The same suite commit has a staging run and a production run in docs/testing/qa-runs, and the register shows each case's status on both sites."),
    ("SC-37", "Signed-in test run by the Claude QA agent", "QA-011, QA-010", "P1",
     "1. The agent signs the staging test accounts in with the emailed codes (owner's mailbox)\n2. It runs the signed-in tier on test.gokesari.com\n3. It collects the results and rebuilds the register",
     "A staging run with signed-in checks that passed (so the sessions worked), collected into docs/testing/qa-runs."),
]

COVERAGE = [
    ("QA-010", "QA", "Automated QA suite: the 791-case plan as UI + API checks (tests/qa), tiers guest / signed / local", "COMPLETED (24-Oct)", "TC-QA-001, TC-QA-002, TC-QA-003, SC-36, SC-37"),
    ("QA-011", "QA", "Claude QA agent and skills (.claude/agents/gokesari-qa.md, .claude/skills/gokesari-qa-*)", "COMPLETED (24-Oct)", "TC-QA-005, SC-37"),
    ("QA-012", "QA", "Master test register with execution status per site (docs/testing/GOKESARI_MASTER_TEST_REGISTER.xlsx)", "COMPLETED (24-Oct)", "TC-QA-004, SC-36"),
    ("QA-013", "QA", "Scheduled QA workflow (weekly + by hand): production-safe tier on test.gokesari.com and gokesari.com; whole suite on a local build by hand", "COMPLETED (24-Oct)", "TC-QA-006, SC-36"),
    ("QA-014", "QA", "Local QA environment: seeded test data, local mail catcher, one-command setup", "COMPLETED (24-Oct)", "TC-QA-007"),
]


LIST_MODULE = "Purchase lists"
LIST_SOURCE = "docs/purchase-lists-2026-10/README.md (24 Oct 2026, PR #231)"

LIST_CASES = [
    # id, features, title, type, priority, role, method, preconditions, steps, expected
    ("TC-LIST-001", "PL-001", "Guest is sent to sign-in from My lists and from a list's page", "Security", "P0", "Guest", "UI",
     "Signed out",
     "1. Open /lists and /lists?kind=MEDICINE\n2. Open /lists/<any id> and /lists/<malformed id>",
     "Each page answers a redirect to /signin; the guest lands on the sign-in page and sees no list."),
    ("TC-LIST-002", "PL-001", "Every purchase-lists API refuses a caller without a session", "Security", "P0", "Guest", "API",
     "Signed out; ids that exist nowhere",
     "1. Call every method of /api/purchase-lists, /:id, /:id/items, /:id/items/:itemId, /:id/check, /:id/order and /from-cart without a session, with a plausible body",
     "Every call answers 401 UNAUTHENTICATED before the body is read: no field errors, no 404 that would tell whether a list exists."),
    ("TC-LIST-003", "PL-001", "Create a weekly shopping list", "Positive", "P0", "Customer", "UI+API",
     "Customer signed in",
     "1. Open My lists (/lists)\n2. In 'Create a list' choose Shopping list, type a name, choose Weekly, leave 'Remind me' on\n3. Create list\n4. Go back to My lists",
     "The new list opens (/lists/<id>) empty, with 'Browse shops'; its header reads 'Shopping list · Weekly'. Saved as SHOPPING / WEEKLY / reminder on, next due 7 days after it was made. My lists shows it under 'Shopping lists' with 'Weekly', 'No products yet' and 'Next due <date>'."),
    ("TC-LIST-004", "PL-001", "Create a monthly medicine list with the reminder off", "Positive", "P1", "Customer", "UI+API",
     "Customer signed in",
     "1. Open /lists?kind=MEDICINE (the form opens on Medicines)\n2. Type a name, choose Monthly, untick 'Remind me when it is due again'\n3. Create list\n4. Go back to My lists",
     "The list opens with 'Medicines · Monthly · … · Reminder off'; the empty list links 'Find a pharmacy' (/shops?type=PHARMACY). Saved as MEDICINE / MONTHLY / reminder off, next due a month after it was made. Shown under 'Medicine lists', not under Shopping lists."),
    ("TC-LIST-005", "PL-001", "A second list with the same name is refused (any case, extra spaces)", "Negative", "P1", "Customer", "UI+API",
     "A list named X exists",
     "1. From the form, create a list named X in other case with extra spaces\n2. POST /api/purchase-lists with X in lower case",
     "The form shows 'You already have a list called …' and stays on My lists; the API answers 409 with the same message; there is still exactly one list named X."),
    ("TC-LIST-006", "PL-002", "Add to list from a product card", "Positive", "P0", "Customer", "UI+API",
     "A list exists; a product of shop A on sale online",
     "1. Open the shop's page and find the product's card\n2. Tap 'Add to list'\n3. Pick the list in the sheet\n4. Tap 'Open'\n5. As a guest, open a shop's page",
     "The sheet 'Add to which list?' shows the customer's lists; picking one adds the product (201) with quantity 1 and the card says 'Added to \"<list>\"' with an Open link; the list shows the product x1. A guest's product cards show no 'Add to list'."),
    ("TC-LIST-007", "PL-002", "Create & add a new list from the product card's sheet", "Positive", "P1", "Customer", "UI+API",
     "A product of shop A on sale online",
     "1. Open the card's 'Add to list' sheet\n2. Under 'Or start a new list' type a name and choose Medicines\n3. Create & add",
     "'Create & add' stays disabled until a name is typed; a new MEDICINE list is made with the product x1 and the card says 'Added to \"<name>\"'."),
    ("TC-LIST-008", "PL-002", "A product already on the list is left as it is ('Already on')", "Negative", "P1", "Customer", "UI+API",
     "The product is on the list x3 with a note",
     "1. Card -> Add to list -> pick that list\n2. POST the same product again with quantity 7 and another note",
     "The card says 'Already on \"<list>\"' (200); the API answers 200 with alreadyOnList; the line keeps quantity 3 and its note; no second line."),
    ("TC-LIST-009", "PL-002", "Only products customers can see can be added", "Security", "P1", "Customer", "API",
     "A list of the customer's",
     "1. POST …/items with a product id no listing has\n2. POST …/items with a malformed id\n3. Add an inactive listing, and a product of a shop that is not approved",
     "Unknown id -> 404 'Product was not found'; malformed id -> 422; an inactive listing or a product of a shop that is not approved -> 404 too, so a list cannot reveal a hidden shop's products. Nothing is added."),
    ("TC-LIST-010", "PL-002", "Quantities: stepper, typed value and skip (0)", "Positive", "P0", "Customer", "UI+API",
     "A list with products P and Q, x1 each",
     "1. Open the list; tap + on P, then type 5 in P's quantity and press Enter\n2. Tap - on Q down to 0\n3. Reload; open My lists",
     "P is saved at 2, then 5; Q at 0 shows 'Skipped this time', its - is disabled and it stays on the list; 'Products to order' reads 1 of 2; the estimate is today's price x quantity of P only. After a reload the values are kept; the list's card reads '2 products · 1 skipped this time'."),
    ("TC-LIST-011", "PL-002", "Notes are saved; quantities 0-99 and notes up to 120 characters", "Boundary", "P1", "Customer", "UI+API",
     "A list with one product",
     "1. Type a note and leave the field; reload\n2. PATCH the item with quantity 100, -1, 1.5 and with a 121-character note\n3. PATCH quantity 99 and a 120-character note",
     "The note is saved when the field is left and kept after a reload; 100, -1, 1.5 and 121 characters are refused (422); 99 and 120 characters are accepted."),
    ("TC-LIST-012", "PL-002", "Remove a product from a list", "Positive", "P1", "Customer", "UI+API",
     "A list with products P and Q",
     "1. Tap Remove on Q\n2. Reload\n3. DELETE the same item again",
     "Q's row disappears (204) and stays gone after a reload; P stays; removing it again answers 404."),
    ("TC-LIST-013", "PL-003", "Check list shows the list priced like the cart", "Positive", "P0", "Customer", "UI+API",
     "Empty cart; a list with P x2 and Q x0",
     "1. Open the list and tap Check list\n2. Read the dialog\n3. Cancel",
     "'Confirm your order' shows, per shop: 'P x 2' at today's price, Subtotal, Delivery (amount or Free) and 'Total to pay' equal to the check API's grand total (subtotal + delivery); 'Skipped this time: Q'; no add / replace choice (the cart is empty); it says nothing is charged yet. Cancel closes it; the cart and the list are unchanged."),
    ("TC-LIST-014", "PL-003", "Check list shows the lines that will be left out, with the reason", "Negative", "P1", "Customer", "UI+API",
     "A list with P x1 and L at more than L's stock",
     "1. Open the list\n2. Tap Check list",
     "The list shows L with 'Only N left in stock.' and '1 product can't be ordered right now and will be left out.'; the dialog shows L's line with the reason and '—' for its price, the same warning, and a total of P alone; Confirm order stays available for P. ('No longer sold' and 'shop not accepting orders' read the same way.)"),
    ("TC-LIST-015", "PL-003", "Check with items already in the cart: add to them or replace the cart", "Positive", "P1", "Customer", "UI",
     "Cart holds 1 item; a list with P x1",
     "1. Tap Check list\n2. Switch from 'Add this list to them' to 'Replace my cart with this list'\n3. Cancel",
     "'Your cart already has 1 item' with the two choices, 'Add this list to them' chosen; the total reads 'Total for this list' with 'Plus what is already in your cart…'; on Replace it reads 'Total to pay' and the note goes. The cart is unchanged."),
    ("TC-LIST-016", "PL-004", "Confirm order with an empty cart", "Positive", "P0", "Customer", "UI+API",
     "Empty cart; a weekly list with P x2 and Q x0",
     "1. Check list -> Confirm order\n2. Read the cart\n3. Open the list again",
     "/cart opens holding P x2 only (the skipped Q is not added); the cart's total equals the checked total; the list shows 'Last ordered <today>' and is next due a week later. No order is placed and nothing is charged."),
    ("TC-LIST-017", "PL-004", "Confirm order: add the list to a cart that has items", "Positive", "P0", "Customer", "UI+API",
     "Cart: Q x1 and P x1; a list with P x2",
     "1. Check list -> keep 'Add this list to them' -> Confirm order",
     "The order call reports P x2 added, nothing left out, 4 units in the cart; /cart opens with Q x1 kept and P x3."),
    ("TC-LIST-018", "PL-004", "Confirm order: replace the cart with the list", "Positive", "P0", "Customer", "UI+API",
     "Cart: Q x1; a list with P x2",
     "1. Check list -> 'Replace my cart with this list' -> Confirm order",
     "/cart opens holding only P x2."),
    ("TC-LIST-019", "PL-004", "Nothing to order is refused", "Negative", "P0", "Customer", "UI+API",
     "A list with every quantity 0; an empty list; cart: Q x1",
     "1. Open the all-skipped list\n2. POST …/order with ADD and with REPLACE; REPLACE for the empty list",
     "Check list is disabled with 'Set a quantity above 0 for something to order.'; each send answers 422 'Nothing to order…'; the cart keeps Q x1; the list is not marked ordered."),
    ("TC-LIST-020", "PL-004", "No line can go in: the cart is left exactly as it was", "Negative", "P1", "Customer", "UI+API",
     "(a) Cart: R x99; a list with R x1. (b) Cart: Q x1; a list with only L, above L's stock",
     "(a) Check list -> Confirm order\n(b) Check list; POST …/order with REPLACE",
     "(a) 'Nothing could be added', 'Your cart was not changed', Left out: 'R — Your cart already has 99 — the most one line can hold.', no 'Go to checkout'; the cart still holds R x99; the list is not marked ordered. (b) Confirm order is disabled; REPLACE answers 200 with nothing added and L left out ('Only N left in stock.'); the cart still holds Q x1 (REPLACE never empties a cart for nothing)."),
    ("TC-LIST-021", "PL-005", "Save the cart as a new list", "Positive", "P1", "Customer", "UI+API",
     "Cart: P x2 and Q x1",
     "1. /cart -> 'Save as a list'\n2. A new list: Shopping list, a name, Weekly -> Save list\n3. Open 'Save as a list' again",
     "'Saved to \"<name>\".' with 'Open the list'; the list is SHOPPING / WEEKLY with P x2 and Q x1; the cart is unchanged; the new list is offered, and chosen, in 'Save to' without a reload."),
    ("TC-LIST-022", "PL-005", "Save the cart into an existing list", "Positive", "P1", "Customer", "UI+API",
     "A list with P x5 (and T x3); cart: P x2 and Q x1",
     "1. /cart -> Save as a list -> Save to: that list -> Save list\n2. Empty the cart and save it into the list again (API)",
     "'Products already on that list take the quantity in your cart.' is shown; the list becomes P x2, Q x1 (T x3 kept); saving an empty cart answers 422 'Your cart is empty'."),
    ("TC-LIST-023", "PL-001", "Edit a list's name, type, frequency and reminder", "Positive", "P1", "Customer", "UI+API",
     "A shopping, monthly list; another list named Y",
     "1. Edit list -> rename to Y (other case) -> Save\n2. Rename to a new name, choose Medicines, Weekly, untick the reminder -> Save\n3. Go back to My lists",
     "Y is refused ('You already have a list called …'); then the header shows the new name and 'Medicines · Weekly · … · Reminder off', next due a week after it was made; the card moves to 'Medicine lists'."),
    ("TC-LIST-024", "PL-001", "Delete a list", "Positive", "P1", "Customer", "UI+API",
     "A list with a product",
     "1. Edit list -> Delete list -> 'Yes, delete list'",
     "It asks 'Delete this list?' once; then My lists opens without it; its page and its API answer 404."),
    ("TC-LIST-025", "PL-001", "Another customer's list is not found", "Security", "P0", "Customer", "UI+API",
     "Customer 1's list with a product; customer 2 signed in",
     "1. As customer 2: GET, PATCH and DELETE the list; POST, PATCH and DELETE its items; GET its check; POST its order\n2. Open /lists/<id> and /lists/not-a-list-id",
     "Every call answers 404 (never 403); both pages answer 404; the list is not among customer 2's lists; customer 1's list is unchanged."),
    ("TC-LIST-026", "PL-006", "Ways into My lists", "Positive", "P1", "Customer", "UI",
     "Customer signed in",
     "1. Account menu -> My lists\n2. Home: the Orders menu -> My lists (account dropdown on wide screens, the menu drawer on phones; the Orders tile on the Tile Board home)\n3. /orders -> the 'My lists' tab",
     "Each opens My lists (/lists)."),
    ("TC-LIST-027", "PL-006", "Phone layout of the lists and the Add to list sheet", "Positive", "P1", "Customer", "UI",
     "Phone (Pixel 7, 412 px); a list with a long name and one product",
     "1. Open My lists, the list and its Check list dialog\n2. On the shop's page open a card's Add to list sheet",
     "No sideways scroll on any of them; the Check list bar sits above the items; Confirm order can be reached; the sheet spans the screen's width, shows the long list name in full, with 44 px tap targets and the new-list field reachable. (Wide screens: no bar, Check list in the side panel.)"),
    ("TC-LIST-028", "PL-007", "Reminder when a list is due again", "Positive", "P2", "System", "DB",
     "A list with something to order, reminder on, last ordered a week (weekly) or a month (monthly) ago",
     "1. Run the daily purchase-list-reminders job (05:00 IST)\n2. Run it again",
     "One in-app 'Time to reorder: <list>' notification linking to the list, once per due date; none for a list with the reminder off, with nothing to order, or for an account that is not active."),
]

LIST_COVERAGE = [
    ("PL-001", "Purchase lists", "Named lists per customer: shopping or medicines, weekly or monthly, reminder on/off; unique names; edit, delete; only the owner sees a list; guests refused", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-001, TC-LIST-002, TC-LIST-003, TC-LIST-004, TC-LIST-005, TC-LIST-023, TC-LIST-024, TC-LIST-025"),
    ("PL-002", "Purchase lists", "Products on a list: 'Add to list' on product cards (pick, create & add, already on), only visible products, quantities 0-99 (0 = skip), notes, remove", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-006, TC-LIST-007, TC-LIST-008, TC-LIST-009, TC-LIST-010, TC-LIST-011, TC-LIST-012"),
    ("PL-003", "Purchase lists", "Check list: the list priced like the cart (per shop, delivery, total), left-out lines with reasons, add-to / replace choice", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-013, TC-LIST-014, TC-LIST-015"),
    ("PL-004", "Purchase lists", "Confirm order: the list to the cart (add or replace), nothing-to-order refusal, the cart never emptied for nothing, list marked ordered", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-016, TC-LIST-017, TC-LIST-018, TC-LIST-019, TC-LIST-020"),
    ("PL-005", "Purchase lists", "'Save as a list' on the cart: a new list or an existing one (cart quantities win)", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-021, TC-LIST-022"),
    ("PL-006", "Purchase lists", "Ways in (account menu, Orders menu, /orders tab) and the phone layout (Check list bar, Add to list sheet)", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-026, TC-LIST-027"),
    ("PL-007", "Purchase lists", "Daily reminder (05:00 IST) for a list that is due again", "BUILT 24-Oct (PR #231, staging pending)",
     "TC-LIST-028"),
]

LIST_CHANGE = ("PL 24-Oct", "New module",
               "Purchase lists (PL-001 to PL-007, docs/purchase-lists-2026-10, PR #231): TC-LIST-001 to TC-LIST-028. Automated in "
               "tests/qa/guest/purchase-lists.spec.ts (guest refusals, production-safe) and tests/qa/signed-in/purchase-lists.spec.ts "
               "(signed tier: staging customer accounts and the local server; writes only 'QA PL' lists and puts the cart back). "
               "TC-LIST-028 (reminders) is time-based and declared manual (integration tests cover it).")

APP_MODULE = "App install & updates"
APP_SOURCE = "docs/app-install-2026-10/README.md (31 Oct 2026)"

APP_CASES = [
    # id, features, title, type, priority, role, method, preconditions, steps, expected
    ("TC-APP-001", "AI-001", "First visit on an Android phone shows the install prompt", "Positive", "P0", "Guest", "UI",
     "Android phone browser, app not installed, fresh browser storage",
     "1. Open the home page",
     "A bottom sheet 'Get the GoKesari app' with Install and Not now appears without a backdrop; the page behind stays usable."),
    ("TC-APP-002", "AI-001", "Not now hides the prompt for this page view only", "Positive", "P0", "Guest", "UI",
     "TC-APP-001 shown",
     "1. Tap Not now\\n2. Read localStorage gk_app_install_declines",
     "The sheet closes; the decline is counted (count 1) but no flag that stops prompting is stored."),
    ("TC-APP-003", "AI-001", "After declining, the prompt shows again on visits 2 and 3", "Positive", "P0", "Guest", "UI",
     "Declined once (TC-APP-002)",
     "1. Open another page (full load) - visit 2\\n2. Tap Not now\\n3. Open the home page again - visit 3",
     "The prompt is shown on visit 2 and on visit 3; the decline count rises to 2. Declining never suppresses it."),
    ("TC-APP-004", "AI-002", "Install from Chrome's own dialog; no prompt after install", "Positive", "P0", "Guest", "UI",
     "Android Chrome fires beforeinstallprompt",
     "1. Tap Install\\n2. Accept Chrome's dialog (appinstalled)\\n3. Reload",
     "Chrome's install dialog opens from the button (its own mini-bar is suppressed); after install the sheet is gone and does not return on reload (gk_app_installed set)."),
    ("TC-APP-005", "AI-002", "Installed app (standalone, installed flag, related app) never shows the prompt", "Positive", "P0", "Guest", "UI",
     "App installed on this device",
     "1. Open the site from the installed app / with the installed flag set",
     "No install prompt anywhere on the website."),
    ("TC-APP-006", "AI-002", "No install prompt inside the native Android / iPhone app", "Positive", "P0", "Guest", "UI",
     "User agent contains GoKesariApp/<version>",
     "1. Open the home page in the app's WebView",
     "No install prompt and no service-worker update bar."),
    ("TC-APP-007", "AI-003", "iPhone Safari shows the Add to Home Screen guide", "Positive", "P0", "Guest", "UI",
     "iPhone Safari, not added to the home screen",
     "1. Open the home page\\n2. Tap Install",
     "The sheet shows; Install opens a three-step picture guide: Tap Share -> Add to Home Screen -> Tap Add."),
    ("TC-APP-008", "AI-004", "Decline and install are kept on the signed-in account", "Positive", "P1", "Customer", "API",
     "Signed in",
     "1. POST /api/app-install declined (x2), installed android\\n2. GET /api/app-install\\n3. Same calls signed out",
     "The record shows declineCount up by 2 and android among installedPlatforms; signed out every call answers 401."),
    ("TC-APP-009", "AI-005", "Service worker served fresh with a versioned cache", "Positive", "P0", "Guest", "API",
     "-",
     "1. GET /sw.js",
     "200 JavaScript, Cache-Control no-store; names its build version and caches only /_next/static and /pwa files in gk-static-<version>."),
    ("TC-APP-010", "AI-005", "A new deploy is picked up by the installed app without reinstalling", "Positive", "P0", "Guest", "UI",
     "App open on the old version; a new version deployed",
     "1. Return to the app / reload\\n2. Tap Reload on 'A new version of GoKesari is ready'",
     "The bar appears; Reload switches to the new version (the old gk-static cache is deleted). Without Reload, the new version runs on the next launch."),
    ("TC-APP-011", "AI-006", "Native update check endpoint", "Positive", "P1", "Guest", "API",
     "-",
     "1. GET /api/app/version",
     "android and ios each with latest, minimum, url (null until configured; versions as x.y.z, links https)."),
    ("TC-APP-012", "AI-006", "Native app offers a newer build and installs it over the old one", "Positive", "P0", "Customer", "Manual",
     "App 1.5.0 installed; MOBILE_APP_<platform>_LATEST set higher with a URL",
     "1. Open the app (or bring it back)\\n2. Tap Update",
     "'A new version of GoKesari is ready' with Later / Update (only Update below the minimum); Update opens the Play Store / App Store / APK and the new build installs over the old one, still signed in."),
    ("TC-APP-013", "AI-001", "The prompt fits a phone screen", "Positive", "P1", "Guest", "UI",
     "Phone 360-412 px wide",
     "1. Open the home page\\n2. Open the guide",
     "The whole sheet (title, text, guide, buttons) is on screen without scrolling, buttons 44 px tall, no sideways scroll, Kesari colours."),
]

APP_COVERAGE = [
    ("AI-001", "App install & updates", "Install prompt on every visit until installed; Not now for the page view only; phone layout (Tile Board)", "BUILT 31-Oct", "TC-APP-001, TC-APP-002, TC-APP-003, TC-APP-013"),
    ("AI-002", "App install & updates", "Install detection: Chrome install dialog, appinstalled, standalone, related apps, native app user agent", "BUILT 31-Oct", "TC-APP-004, TC-APP-005, TC-APP-006"),
    ("AI-003", "App install & updates", "iPhone / iPad: Add to Home Screen guide; store links when listed", "BUILT 31-Oct", "TC-APP-007"),
    ("AI-004", "App install & updates", "Server-side record for signed-in people (app_install_states)", "BUILT 31-Oct", "TC-APP-008"),
    ("AI-005", "App install & updates", "Service worker with versioned cache; Reload bar; update on next launch", "BUILT 31-Oct", "TC-APP-009, TC-APP-010"),
    ("AI-006", "App install & updates", "Native app update check (Android / iPhone 1.5.0)", "BUILT 31-Oct", "TC-APP-011, TC-APP-012"),
]

APP_CHANGE = ("AI 31-Oct", "New module",
              "App install & updates (AI-001 to AI-006, docs/app-install-2026-10): TC-APP-001 to TC-APP-013. Automated in "
              "tests/qa/guest/app-install.spec.ts (production-safe, Android and iPhone profiles) and tests/qa/signed-in/app-install.spec.ts "
              "(account record). TC-APP-010 needs a deploy and TC-APP-012 a phone: declared manual.")


def copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format


def append_like(ws, values, template_row):
    row = ws.max_row + 1
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        copy_style(ws.cell(row=template_row, column=col), cell)
    return row


def ids_in(ws):
    return {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}


def stretch_status_colours(ws, column, last):
    """The Status colours (Pass / Fail / Blocked) reach the last case row."""
    rebuilt = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        ref = str(cf.sqref)
        if re.fullmatch(rf"{column}2:{column}\d+", ref):
            ref = f"{column}2:{column}{last}"
        for rule in cf.rules:
            rebuilt.add(ref, rule)
    ws.conditional_formatting = rebuilt


def refresh_readme(wb):
    """'N cases across M modules (P0 a, P1 b, P2 c; …)' counted from the Test cases sheet."""
    tc = wb["Test cases"]
    rows = [(tc.cell(row=r, column=2).value, tc.cell(row=r, column=6).value) for r in range(2, tc.max_row + 1) if tc.cell(row=r, column=1).value]
    modules = len({m for m, _ in rows})
    by_priority = {p: sum(1 for _, x in rows if x == p) for p in ("P0", "P1", "P2")}
    readme = wb["Read me"]
    for r in range(1, readme.max_row + 1):
        v = readme.cell(row=r, column=2).value
        if isinstance(v, str) and re.match(r"^\d+ cases across \d+ modules \(P0 \d+, P1 \d+, P2 \d+", v):
            v = re.sub(r"^\d+ cases across \d+ modules \(P0 \d+, P1 \d+, P2 \d+",
                       f"{len(rows)} cases across {modules} modules (P0 {by_priority['P0']}, P1 {by_priority['P1']}, P2 {by_priority['P2']}", v)
            readme.cell(row=r, column=2, value=v)
    return readme


def add_module(module, cases, source, scenarios=(), coverage=(), change=None):
    """Adds a module's cases (and scenarios, coverage, change-log line) once; existing rows are left alone."""
    wb = openpyxl.load_workbook(PLAN)
    tc = wb["Test cases"]
    have = ids_in(tc)
    template = tc.max_row
    added = 0
    for case in cases:
        if case[0] in have:
            continue
        cid, features, title, typ, prio, role, method, pre, steps, expected = case
        append_like(tc, [cid, module, features, title, typ, prio, role, method, pre, steps, expected,
                         source, "Not run", None, None, None], template)
        added += 1
    last = tc.max_row
    # Status drop-down, its colours and the filter reach the new rows.
    for dv in tc.data_validations.dataValidation:
        if "M2:M" in str(dv.sqref):
            dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"M2:M{last}")
    stretch_status_colours(tc, "M", last)
    tc.auto_filter.ref = f"A1:P{last}"

    sc = wb["E2E scenarios"]
    have = ids_in(sc)
    template = sc.max_row
    for s in scenarios:
        if s[0] in have:
            continue
        append_like(sc, [s[0], s[1], s[2], s[3], s[4], s[5], "Not run", f"Added 24-Oct ({module})"], template)
        added += 1
    for dv in sc.data_validations.dataValidation:
        m = re.match(r"G2:G\d+", str(dv.sqref))
        if m:
            dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"G2:G{sc.max_row}")
    if sc.auto_filter.ref:
        sc.auto_filter.ref = re.sub(r"\d+$", str(sc.max_row), sc.auto_filter.ref)

    # Run summary: ranges reach the last case row; a row for the new module before Total.
    rs = wb["Run summary"]
    total_row = next(r for r in range(2, rs.max_row + 1) if rs.cell(row=r, column=1).value == "Total")
    modules = [rs.cell(row=r, column=1).value for r in range(2, total_row)]
    if module not in modules:
        rs.insert_rows(total_row)
        for col in range(1, 9):
            copy_style(rs.cell(row=total_row - 1, column=col), rs.cell(row=total_row, column=col))
        rs.cell(row=total_row, column=1, value=module)
        total_row += 1
    # Every module row and the Total row, written in the sheet's own formula pattern over all case rows.
    rng = lambda c: f"'Test cases'!${c}$2:${c}${last}"
    for r in range(2, total_row):
        rs.cell(row=r, column=2, value=f"=COUNTIF({rng('B')},A{r})")
        for col, status in zip(range(3, 7), ["Pass", "Fail", "Blocked", "Not run"]):
            rs.cell(row=r, column=col, value=f'=COUNTIFS({rng("B")},$A{r},{rng("M")},"{status}")')
        rs.cell(row=r, column=7, value=f'=COUNTIFS({rng("B")},$A{r},{rng("F")},"P0")')
        rs.cell(row=r, column=8, value=f'=G{r}-COUNTIFS({rng("B")},$A{r},{rng("F")},"P0",{rng("M")},"Pass")')
    for col in range(2, 9):
        letter = chr(64 + col)
        rs.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{total_row - 1})")

    cov = wb["Coverage"]
    have = ids_in(cov)
    template = cov.max_row
    for fid, mod, feature, status, tcs in coverage:
        if fid in have:
            continue
        append_like(cov, [fid, mod, feature, status, len(tcs.split(", ")), tcs, ""], template)
    if cov.auto_filter.ref:
        cov.auto_filter.ref = re.sub(r"\d+$", str(cov.max_row), cov.auto_filter.ref)

    if change:
        log = wb["Change log"]
        if not any(log.cell(row=r, column=1).value == change[0] for r in range(2, log.max_row + 1)):
            append_like(log, list(change), log.max_row)
        if log.auto_filter.ref:
            log.auto_filter.ref = re.sub(r"\d+$", str(log.max_row), log.auto_filter.ref)
    wb.save(PLAN)
    print(f"plan_workbook: {added} rows added to '{module}'; Test cases now {last - 1} cases, scenarios {sc.max_row - 1}")
    return wb


def add_qa_cases():
    add_module(MODULE, QA_CASES, "QA automation, 24 Oct 2026 (tests/qa/README.md)", QA_SCENARIOS, COVERAGE,
               ("QA-AUTO 24-Oct", "New module",
                "QA automation (QA-010 to QA-014): TC-QA-001 to TC-QA-007 and scenarios SC-36, SC-37. "
                "Status, Actual result, Defect / evidence and Run date are now written from the automated runs "
                "(docs/testing/GOKESARI_MASTER_TEST_REGISTER.xlsx is the full record per site)."))
    wb = openpyxl.load_workbook(PLAN)
    readme = wb["Read me"]
    for r in range(1, readme.max_row + 1):
        v = readme.cell(row=r, column=2).value
        if isinstance(v, str) and v.startswith("791 cases across 28 modules (P0 306, P1 325, P2 160)."):
            readme.cell(row=r, column=2, value=v.replace("791 cases across 28 modules (P0 306, P1 325, P2 160).",
                                                         "798 cases across 29 modules (P0 309, P1 329, P2 160; the QA automation module TC-QA-001 to 007 added 24 Oct)."))
    readme = refresh_readme(wb)
    for r in range(1, readme.max_row + 1):
        v = readme.cell(row=r, column=2).value
        if isinstance(v, str) and v.startswith("35 end-to-end journeys") and "SC-36" not in v:
            readme.cell(row=r, column=2, value=v.replace("35 end-to-end journeys", "37 end-to-end journeys").rstrip(".") +
                        "; SC-36 and SC-37 (24 Oct) are the QA release check and the Claude QA agent's signed-in run.")
    wb.save(PLAN)


def add_list_cases():
    add_module(LIST_MODULE, LIST_CASES, LIST_SOURCE, (), LIST_COVERAGE, LIST_CHANGE)
    wb = openpyxl.load_workbook(PLAN)
    readme = refresh_readme(wb)
    for r in range(1, readme.max_row + 1):
        v = readme.cell(row=r, column=2).value
        if isinstance(v, str) and "TC-QA-001 to 007 added 24 Oct" in v and "TC-LIST-001" not in v:
            readme.cell(row=r, column=2, value=v.replace("TC-QA-001 to 007 added 24 Oct", "TC-QA-001 to 007 and the Purchase lists module TC-LIST-001 to 028 added 24 Oct"))
    wb.save(PLAN)
    print("plan_workbook: Read me counts refreshed")


STATUS_MAP = {"Pass": "Pass", "Pass (partial)": "Pass", "Fail": "Fail", "Blocked": "Blocked", "Manual": "Blocked",
              "Not run (rule)": "Not run", "Not run": "Not run"}


def statuses(register):
    reg = openpyxl.load_workbook(register, read_only=True)
    rows = {}
    for sheet in ("Test cases", "E2E scenarios"):
        ws = reg[sheet]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(header)}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0]:
                rows[r[0]] = {"overall": r[idx["Overall"]], "detail": r[idx["Result detail (latest, per site)"]] or "",
                              "defect": r[idx["Defect"]] or "", "last": r[idx["Last run (UTC)"]] or "",
                              "local": r[idx["Local"]], "staging": r[idx["Staging (test.gokesari.com)"]],
                              "production": r[idx["Production (gokesari.com)"]]}
    wb = openpyxl.load_workbook(PLAN)
    tc = wb["Test cases"]
    n = 0
    for r in range(2, tc.max_row + 1):
        cid = tc.cell(row=r, column=1).value
        if cid not in rows:
            continue
        x = rows[cid]
        tc.cell(row=r, column=13, value=STATUS_MAP.get(x["overall"], "Not run"))
        sites = f"Local {x['local']} · Staging {x['staging']} · Production {x['production']}"
        prefix = "Passed, partly covered. " if x["overall"] == "Pass (partial)" else "Manual (not automatable): " if x["overall"] == "Manual" else ""
        tc.cell(row=r, column=14, value=f"{prefix}{sites}\n{x['detail']}"[:1500] if x["overall"] != "Not run" else None)
        tc.cell(row=r, column=15, value=x["defect"] or None)
        tc.cell(row=r, column=16, value=x["last"] or None)
        n += 1
    sc = wb["E2E scenarios"]
    for r in range(2, sc.max_row + 1):
        sid = sc.cell(row=r, column=1).value
        if sid not in rows:
            continue
        x = rows[sid]
        sc.cell(row=r, column=7, value=STATUS_MAP.get(x["overall"], "Not run"))
        note = sc.cell(row=r, column=8).value or ""
        note = re.sub(r"\n?\[Run .*\]$", "", note, flags=re.S)
        sc.cell(row=r, column=8, value=f"{note}\n[Run {x['last']}: Local {x['local']} · Staging {x['staging']} · Production {x['production']}. {x['detail'][:600]}]".strip())
        n += 1
    wb.save(PLAN)
    print(f"plan_workbook: statuses written for {n} cases and scenarios")


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "add-qa-cases":
        add_qa_cases()
    elif len(sys.argv) >= 2 and sys.argv[1] == "add-list-cases":
        add_list_cases()
    elif len(sys.argv) >= 2 and sys.argv[1] == "add-app-cases":
        add_module(APP_MODULE, APP_CASES, APP_SOURCE, (), APP_COVERAGE, APP_CHANGE)
    elif len(sys.argv) >= 3 and sys.argv[1] == "statuses":
        statuses(sys.argv[2])
    else:
        print(__doc__)
        sys.exit(2)
