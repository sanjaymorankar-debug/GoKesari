"""Batch 2026-10 — updates the product tracker (feature status report).

    python3 scripts/reports/batch-2026-10-tracker.py <base.xlsx> <statuses.json> <out.xlsx>

<base.xlsx> is the latest dated GOKESARI_FEATURE_STATUS_REPORT_*.xlsx. The
"Process Flow Tree" sheet keeps its existing Status column (COMPLETED / IN
PROGRESS / YET TO START / BLOCKED) and its roll-up formulas; this script adds
three columns for the batch's vocabulary —

    Batch status (18 Oct) | Batch PR(s) | Batch gap / note

with DONE / DONE-SANDBOX / DROPPED / DEFERRED-PHASE-2 / ON-HOLD /
BLOCKED-OWNER-INPUT, and parent roll-ups ("done / total" for every level-1
and level-2 row: DONE and DONE-SANDBOX count as done). The existing Status
column is also moved on for the batch items (DONE/DONE-SANDBOX -> COMPLETED,
DROPPED/DEFERRED/ON-HOLD/BLOCKED-OWNER-INPUT -> BLOCKED) so the old roll-ups
stay meaningful. A "Batch 2026-10" sheet lists every batch step.

<statuses.json>: {"F1.1.2": {"status": "DONE", "prs": "#190", "note": "..."}, ...}
"""
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

BATCH = ["DONE", "DONE-SANDBOX", "DROPPED", "DEFERRED-PHASE-2", "ON-HOLD", "BLOCKED-OWNER-INPUT"]
TO_LEGACY = {
    "DONE": "COMPLETED",
    "DONE-SANDBOX": "COMPLETED",
    "DROPPED": "BLOCKED",
    "DEFERRED-PHASE-2": "BLOCKED",
    "ON-HOLD": "BLOCKED",
    "BLOCKED-OWNER-INPUT": "BLOCKED",
}
FILL = {
    "DONE": "C6EFCE",
    "DONE-SANDBOX": "DDEBF7",
    "DROPPED": "D9D9D9",
    "DEFERRED-PHASE-2": "FFF2CC",
    "ON-HOLD": "FFF2CC",
    "BLOCKED-OWNER-INPUT": "F8CBAD",
}


def main(base: str, statuses_path: str, out: str) -> None:
    statuses = json.load(open(statuses_path))
    wb = openpyxl.load_workbook(base)
    ws = wb["Process Flow Tree"]
    header_row = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Flow / step ID")
    first, last = header_row + 1, ws.max_row
    col_status, col_pr, col_note = 10, 11, 12  # J, K, L
    for col, title in ((col_status, "Batch status (18 Oct)"), (col_pr, "Batch PR(s)"), (col_note, "Batch gap / note")):
        cell = ws.cell(header_row, col, title)
        cell.font = Font(bold=True)
    ws.column_dimensions["J"].width = 22
    ws.column_dimensions["K"].width = 14
    ws.column_dimensions["L"].width = 70

    seen = set()
    for r in range(first, last + 1):
        step = ws.cell(r, 1).value
        if step in statuses:
            s = statuses[step]
            assert s["status"] in BATCH, f"{step}: unknown status {s['status']}"
            ws.cell(r, col_status, s["status"]).fill = PatternFill("solid", fgColor=FILL[s["status"]])
            ws.cell(r, col_pr, s.get("prs", ""))
            ws.cell(r, col_note, s.get("note", "")).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, 7, TO_LEGACY[s["status"]])
            # The old note ("Waits on D2 …") is superseded for batch steps.
            ws.cell(r, 9, f'{s["status"]} (18 Oct batch): {s.get("note", "")}')
            prs = s.get("prs", "")
            old = ws.cell(r, 8).value
            if prs and (not old or old == "-"):
                ws.cell(r, 8, prs)
            elif prs and prs not in str(old):
                ws.cell(r, 8, f"{old}; {prs}")
            seen.add(step)
    missing = sorted(set(statuses) - seen)
    assert not missing, f"steps not found in the tracker: {missing}"

    # Roll-ups for the batch column: every level-1/level-2 row counts its level-3 children.
    rows = [(r, str(ws.cell(r, 2).value)) for r in range(first, last + 1)]
    for i, (r, level) in enumerate(rows):
        if level not in ("1", "2"):
            continue
        end = last
        for r2, level2 in rows[i + 1 :]:
            if level2 == "1" or (level == "2" and level2 == "2"):
                end = r2 - 1
                break
        rng_b, rng_j = f"$B${r + 1}:$B${end}", f"$J${r + 1}:$J${end}"
        ws.cell(r, col_status).value = (
            f'=(COUNTIFS({rng_b},3,{rng_j},"DONE")+COUNTIFS({rng_b},3,{rng_j},"DONE-SANDBOX"))'
            f'&" / "&COUNTIFS({rng_b},3,{rng_j},"<>")&" batch steps done"'
        )
        ws.cell(r, col_status).font = Font(bold=True)

    # One sheet listing the batch.
    if "Batch 2026-10" in wb.sheetnames:
        del wb["Batch 2026-10"]
    bs = wb.create_sheet("Batch 2026-10", 1)
    bs.append(["GoKesari build & integrate batch (18 Oct 2026) — status of every flow step in the brief"])
    bs.append(["DONE = built, merged, tested; DONE-SANDBOX = complete on the sandbox/mock provider, production flag OFF until the owner's keys exist"])
    bs.append([])
    bs.append(["Flow / step ID", "Step", "Batch status", "PR(s)", "Gap / note"])
    for c in bs[4]:
        c.font = Font(bold=True)
    for r in range(first, last + 1):
        step = ws.cell(r, 1).value
        if step in statuses:
            s = statuses[step]
            bs.append([step, ws.cell(r, 3).value, s["status"], s.get("prs", ""), s.get("note", "")])
            bs.cell(bs.max_row, 3).fill = PatternFill("solid", fgColor=FILL[s["status"]])
    bs.append([])
    bs.append(["Totals"])
    for status in BATCH:
        bs.append([status, f'=COUNTIF($C$5:$C${4 + len(statuses)},"{status}")'])
    for col, width in zip("ABCDE", (14, 50, 22, 12, 90)):
        bs.column_dimensions[col].width = width
    for row in bs.iter_rows(min_row=5, max_row=4 + len(statuses)):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out)
    print(f"updated {len(seen)} steps -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
