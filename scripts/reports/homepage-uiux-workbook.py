#!/usr/bin/env python3
"""
Homepage UI/UX upgrade (docs/homepage-uiux-2026-10): the status workbook —
product tracker update and testing report in one file.

    python3 scripts/reports/homepage-uiux-workbook.py <out.xlsx>

Reads the run log (docs/homepage-uiux-2026-10/run-log.json: deploy times,
Lighthouse, suite counts, regression results, defects, decisions) and the
Playwright result files the homepage spec writes
(docs/homepage-uiux-2026-10/results/<host>[-app-<os>].json), one per environment.
Sheets: Summary, Features, Flows, Test Cases, Regression, Defects,
Lighthouse, Decisions — colour-coded status cells, frozen headers, filters.
"""
import json
import os
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RUN_LOG = os.path.join(ROOT, "docs/homepage-uiux-2026-10/run-log.json")

ENVS = OrderedDict(
    [
        ("Local", "localhost_3000"),
        ("Staging", "test.gokesari.com"),
        ("Production", "gokesari.com"),
        ("Android", "localhost_3000-app-android"),
        ("iOS", "localhost_3000-app-ios"),
    ]
)

GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="C2410C")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {
    "done": GREEN,
    "pass": GREEN,
    "passed": GREEN,
    "fixed": GREEN,
    "closed": GREEN,
    "in progress": AMBER,
    "partial": AMBER,
    "open": AMBER,
    "pending": AMBER,
    "skipped": GREY,
    "n/a": GREY,
    "not run": GREY,
    "not started": GREY,
    "fail": RED,
    "failed": RED,
    "blocked": RED,
}


def fill_for(value):
    if not isinstance(value, str):
        return None
    key = value.strip().lower()
    if key in STATUS_FILL:
        return STATUS_FILL[key]
    for word, fill in (("blocked", RED), ("fail", RED), ("pass", GREEN), ("done", GREEN), ("fixed", GREEN), ("progress", AMBER), ("emulat", AMBER)):
        if key.startswith(word):
            return fill
    return None


def load(path, default):
    try:
        with open(path, encoding="utf-8") as fh:
            return json.load(fh)
    except FileNotFoundError:
        return default


def results():
    """{env: {case id: row}} from the Playwright result files."""
    out = {}
    for env, run in ENVS.items():
        data = load(os.path.join(ROOT, f"docs/homepage-uiux-2026-10/results/{run}.json"), {"rows": []})
        out[env] = {r["id"]: r for r in data.get("rows", [])}
    return out


def sheet(wb, title, headers, rows, widths=None, status_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if headers[cell.column - 1] in status_cols or cell.column - 1 in status_cols:
                f = fill_for(cell.value)
                if f:
                    cell.fill = f
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(header) or min(max(12, len(header) + 4), 40)
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def main(out_path):
    log = load(RUN_LOG, {})
    res = results()
    env_names = list(ENVS.keys())

    def env_status(case_ids):
        """Pass / Fail / Not run per environment for a list of case ids (all must pass)."""
        statuses = []
        for env in env_names:
            rows = [res[env].get(c) for c in case_ids]
            rows = [r for r in rows if r]
            if not rows:
                statuses.append(log.get("notRunReason", {}).get(env, "Not run"))
            elif all(r["status"] == "pass" for r in rows):
                statuses.append("Pass")
            else:
                statuses.append("Fail")
        return statuses

    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- Summary
    features = log.get("features", [])
    flows = log.get("flows", [])
    all_cases = sorted({cid for env in env_names for cid in res[env]})
    summary_rows = [
        ["Report", "GoKesari — Homepage UI/UX upgrade (F1–F6): status, product tracker update and testing report"],
        ["Generated", log.get("generatedAt", "")],
        ["Branch / PR", log.get("pr", "")],
        ["Features (sub-items)", len(features)],
        ["Features done", sum(1 for f in features if f[4] == "Done")],
        ["End-to-end flows", len(flows)],
        ["Automated test cases (Playwright, unique)", len(all_cases)],
    ]
    for env in env_names:
        rows = list(res[env].values())
        passed = sum(1 for r in rows if r["status"] == "pass")
        rate = f"{passed}/{len(rows)} ({round(100 * passed / len(rows))}%)" if rows else log.get("notRunReason", {}).get(env, "Not run")
        summary_rows.append([f"Playwright pass rate — {env}", rate])
    for label, value in log.get("suites", {}).items():
        summary_rows.append([label, value])
    for label, value in log.get("deploys", {}).items():
        summary_rows.append([f"Deploy — {label}", value])
    blockers = [d for d in log.get("defects", []) if str(d[5]).lower().startswith("blocked")]
    summary_rows.append(["Open blockers", len(blockers) + len(log.get("blockers", []))])
    for b in log.get("blockers", []):
        summary_rows.append(["Blocker", b])
    ws = sheet(wb, "Summary", ["Item", "Value"], summary_rows, {"Item": 44, "Value": 110})
    for r in ws.iter_rows(min_row=2):
        f = fill_for(r[1].value) if isinstance(r[1].value, str) else None
        if r[0].value == "Blocker":
            r[1].fill = RED
        elif f:
            r[1].fill = f

    # --------------------------------------------------------------- Features
    sheet(
        wb,
        "Features",
        ["ID", "Feature", "Sub-item", "Files changed", "Status", *env_names, "Notes"],
        [[*f[:5], *(f[5] if len(f) > 5 and f[5] else env_status(f[6] if len(f) > 6 else [])), f[7] if len(f) > 7 else ""] for f in features],
        {"Feature": 22, "Sub-item": 46, "Files changed": 46, "Notes": 50},
        status_cols=("Status", *env_names),
    )

    # ------------------------------------------------------------------ Flows
    sheet(
        wb,
        "Flows",
        ["Flow ID", "Flow", "New / existing", "Steps", "Expected result", *env_names, "Evidence / notes"],
        [[fl["id"], fl["name"], fl["kind"], fl["steps"], fl["expected"], *(fl.get("results") or env_status(fl.get("cases", []))), fl.get("notes", "")] for fl in flows],
        {"Flow": 30, "Steps": 50, "Expected result": 44, "Evidence / notes": 50},
        status_cols=env_names,
    )

    # ------------------------------------------------------------- Test Cases
    tc_rows = []
    for cid in all_cases:
        sample = next(res[env][cid] for env in env_names if cid in res[env])
        per_env = []
        evidence = []
        for env in env_names:
            row = res[env].get(cid)
            per_env.append(("Pass" if row["status"] == "pass" else f"Fail: {row['detail'][:120]}") if row else "Not run")
            if row and row.get("evidence"):
                evidence.append(f"{env}: {row['evidence']}")
        tc_rows.append([cid, sample["feature"], sample["scenario"], sample["device"], sample["lang"], sample["user"], *per_env, "\n".join(evidence)])
    for extra in log.get("extraTestCases", []):
        tc_rows.append(extra)
    sheet(
        wb,
        "Test Cases",
        ["Test ID", "Feature", "Scenario", "Device / breakpoint", "Language", "User type", *env_names, "Evidence (screenshot)"],
        tc_rows,
        {"Scenario": 40, "Evidence (screenshot)": 70},
        status_cols=env_names,
    )

    # ------------------------------------------------------------- Regression
    sheet(
        wb,
        "Regression",
        ["Existing feature", "How it was checked", "Local", "Staging", "Production", "Notes"],
        log.get("regression", []),
        {"Existing feature": 34, "How it was checked": 60, "Notes": 60},
        status_cols=("Local", "Staging", "Production"),
    )

    # ---------------------------------------------------------------- Defects
    sheet(
        wb,
        "Defects",
        ["ID", "Description", "Severity", "Environment found", "Fix", "Status"],
        log.get("defects", []),
        {"Description": 60, "Fix": 60},
        status_cols=("Status",),
    )

    # ------------------------------------------------------------- Lighthouse
    sheet(
        wb,
        "Lighthouse",
        ["Environment", "When", "Page", "Form factor", "Performance", "Accessibility", "Best practices", "SEO", "Target (Perf ≥ 85, A11y ≥ 95)", "Notes"],
        log.get("lighthouse", []),
        {"Notes": 60, "Target (Perf ≥ 85, A11y ≥ 95)": 22},
        status_cols=("Target (Perf ≥ 85, A11y ≥ 95)",),
    )

    # -------------------------------------------------------------- Decisions
    sheet(
        wb,
        "Decisions",
        ["ID", "Decision", "Why", "Where"],
        log.get("decisions", []),
        {"Decision": 60, "Why": 60, "Where": 40},
    )

    wb.save(out_path)
    print(f"wrote {out_path}: {len(all_cases)} test cases, {len(features)} features, {len(flows)} flows")


if __name__ == "__main__":
    main(sys.argv[1])
